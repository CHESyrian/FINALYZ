"""
Export analysis results to PDF and Excel.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
    ListFlowable,
    ListItem,
)


def build_excel(report: dict[str, Any]) -> bytes:
    """Build an .xlsx workbook from a stored report dict."""
    wb = Workbook()

    # --- Summary sheet ---
    ws = wb.active
    ws.title = "Summary"

    header_font = Font(bold=True, size=14, color="FFFFFF")
    sub_font = Font(bold=True, size=11)
    header_fill = PatternFill("solid", fgColor="0D6EFD")
    good_fill = PatternFill("solid", fgColor="D1E7DD")
    warn_fill = PatternFill("solid", fgColor="FFF3CD")
    bad_fill = PatternFill("solid", fgColor="F8D7DA")
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    ws["A1"] = "Financial Analysis Report"
    ws["A1"].font = header_font
    ws["A1"].fill = header_fill
    ws.merge_cells("A1:B1")

    meta = [
        ("Company", report.get("company_name") or "—"),
        ("Period", report.get("period") or "—"),
        ("Industry", report.get("industry_label") or report.get("industry") or "—"),
        ("Summary", report.get("summary") or "—"),
    ]
    scorecard = report.get("scorecard") or {}
    if scorecard.get("available"):
        meta.append((
            "Industry scorecard",
            f"Grade {scorecard.get('grade')} — {scorecard.get('summary') or ''}",
        ))
    row = 3
    for label, value in meta:
        ws.cell(row, 1, label).font = sub_font
        ws.cell(row, 2, value)
        row += 1

    row += 1
    for title, key in (
        ("Strengths", "strengths"),
        ("Weaknesses", "weaknesses"),
        ("Recommendations", "recommendations"),
    ):
        ws.cell(row, 1, title).font = sub_font
        row += 1
        items = report.get(key) or []
        if not items:
            ws.cell(row, 1, "—")
            row += 1
        for item in items:
            ws.cell(row, 1, str(item))
            row += 1
        row += 1

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 80

    # --- Ratios sheet ---
    ws2 = wb.create_sheet("Ratios")
    headers = [
        "Category", "Ratio", "Abbr", "Value", "Unit", "Status",
        "Industry median", "vs Benchmark", "Interpretation", "Formula",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws2.cell(1, col, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    status_fills = {"good": good_fill, "warning": warn_fill, "bad": bad_fill}

    for r_idx, ratio in enumerate(report.get("ratios") or [], 2):
        values = [
            ratio.get("category", ""),
            ratio.get("name", ""),
            ratio.get("abbreviation", ""),
            ratio.get("formatted", ""),
            ratio.get("unit", ""),
            ratio.get("status", ""),
            ratio.get("benchmark_median") or "—",
            ratio.get("benchmark_label") or "—",
            ratio.get("interpretation", ""),
            ratio.get("formula", ""),
        ]
        for c, v in enumerate(values, 1):
            cell = ws2.cell(r_idx, c, v)
            cell.border = thin
            if c == 6:
                cell.fill = status_fills.get(str(v), PatternFill())

    widths = [14, 28, 8, 12, 8, 10, 16, 24, 45, 40]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # --- DuPont sheet ---
    dupont = report.get("dupont") or {}
    if dupont.get("available"):
        ws3 = wb.create_sheet("DuPont")
        ws3["A1"] = f"DuPont Analysis ({dupont.get('model') or ''})"
        ws3["A1"].font = header_font
        ws3["A1"].fill = header_fill
        ws3.merge_cells("A1:D1")
        ws3["A3"] = "Direct ROE"
        ws3["B3"] = dupont.get("roe_direct_formatted") or "—"
        ws3["A4"] = "Factors product"
        ws3["B4"] = dupont.get("product_formatted") or "—"
        ws3["A5"] = "Narrative"
        ws3["B5"] = dupont.get("narrative") or "—"
        row = 7
        for h, col in (("Component", 1), ("Abbr", 2), ("Value", 3), ("Role", 4), ("Formula", 5)):
            cell = ws3.cell(row, col, h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for c in dupont.get("components") or []:
            row += 1
            ws3.cell(row, 1, c.get("name") or "")
            ws3.cell(row, 2, c.get("abbreviation") or "")
            ws3.cell(row, 3, c.get("formatted") or "")
            ws3.cell(row, 4, c.get("role") or "")
            ws3.cell(row, 5, c.get("formula") or "")
        row += 2
        ws3.cell(row, 1, "Driver notes").font = sub_font
        for note in dupont.get("drivers") or []:
            row += 1
            ws3.cell(row, 1, str(note))
        for i, w in enumerate([22, 10, 12, 40, 40], 1):
            ws3.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(report: dict[str, Any]) -> bytes:
    """Build a PDF report from a stored report dict."""
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title="Financial Analysis Report",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleFA",
        parent=styles["Heading1"],
        fontSize=16,
        spaceAfter=6,
        textColor=colors.HexColor("#0d6efd"),
    )
    h2 = ParagraphStyle(
        "H2FA",
        parent=styles["Heading2"],
        fontSize=12,
        spaceBefore=10,
        spaceAfter=4,
    )
    body = ParagraphStyle(
        "BodyFA",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
    )
    small = ParagraphStyle(
        "SmallFA",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#555555"),
    )

    story: list[Any] = []

    company = report.get("company_name") or "Company"
    period = report.get("period") or ""
    industry = report.get("industry_label") or report.get("industry") or "General"

    story.append(Paragraph("Financial Analysis Report", title_style))
    story.append(Paragraph(
        f"<b>{company}</b>"
        + (f" &nbsp;|&nbsp; {period}" if period else "")
        + f" &nbsp;|&nbsp; Industry: {industry}",
        body,
    ))
    story.append(Spacer(1, 6))

    story.append(Paragraph("Summary", h2))
    story.append(Paragraph(report.get("summary") or "—", body))

    scorecard = report.get("scorecard") or {}
    if scorecard.get("available"):
        story.append(Paragraph("Industry benchmark scorecard", h2))
        story.append(Paragraph(
            f"<b>Grade {scorecard.get('grade')}</b> — {scorecard.get('grade_label') or ''}",
            body,
        ))
        story.append(Paragraph(str(scorecard.get("summary") or ""), small))

    def _bullet_section(title: str, items: list[str]) -> None:
        story.append(Paragraph(title, h2))
        if not items:
            story.append(Paragraph("None identified.", small))
            return
        story.append(
            ListFlowable(
                [ListItem(Paragraph(str(i), small), leftIndent=8) for i in items],
                bulletType="bullet",
                start="•",
            )
        )

    _bullet_section("Strengths", report.get("strengths") or [])
    _bullet_section("Weaknesses", report.get("weaknesses") or [])
    _bullet_section("Recommendations", report.get("recommendations") or [])

    dupont = report.get("dupont") or {}
    if dupont.get("available"):
        story.append(Paragraph(
            f"DuPont Analysis ({dupont.get('model') or ''})",
            h2,
        ))
        story.append(Paragraph(
            f"Direct ROE: <b>{dupont.get('roe_direct_formatted') or '—'}</b>"
            f" &nbsp;|&nbsp; Factors product: <b>{dupont.get('product_formatted') or '—'}</b>",
            body,
        ))
        if dupont.get("narrative"):
            story.append(Paragraph(str(dupont["narrative"]), small))
        d_header = ["Abbr", "Component", "Value", "Role"]
        d_data = [d_header]
        for c in dupont.get("components") or []:
            d_data.append([
                c.get("abbreviation") or "",
                c.get("name") or "",
                c.get("formatted") or "",
                Paragraph(str(c.get("role") or ""), small),
            ])
        d_table = Table(d_data, colWidths=[18 * mm, 42 * mm, 22 * mm, 70 * mm])
        d_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d6efd")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(Spacer(1, 4))
        story.append(d_table)
        if dupont.get("drivers"):
            _bullet_section("DuPont driver notes", list(dupont["drivers"]))

    story.append(Paragraph("Ratios", h2))

    header = ["Abbr", "Ratio", "Value", "Status", "Industry median", "vs Bench."]
    data = [header]
    for r in report.get("ratios") or []:
        data.append([
            r.get("abbreviation") or "",
            Paragraph(r.get("name") or "", small),
            r.get("formatted") or "N/A",
            (r.get("status") or "").upper(),
            r.get("benchmark_median") or "—",
            Paragraph(r.get("benchmark_label") or "—", small),
        ])

    table = Table(data, colWidths=[28, 110, 48, 48, 70, 100])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d6efd")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("FONTSIZE", (0, 1), (-1, -1), 8),
        ("ALIGN", (2, 0), (3, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fa")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(table)

    story.append(Spacer(1, 10))
    story.append(Paragraph(
        "Benchmarks are approximate educational references by industry and are not investment advice.",
        small,
    ))

    doc.build(story)
    return buf.getvalue()
